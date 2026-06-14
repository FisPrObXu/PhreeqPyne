from __future__ import annotations

import csv
import json
import math
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs"
ASSET_DIR = OUT_DIR / "paper_assets"
DOCX_PATH = OUT_DIR / "CFAS_hydrothermal_modeling_paper.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(95, 95, 95)
GRID = "D9E2F3"
HEADER_FILL = "F4F6F9"


def read_rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", newline="", encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


def fval(row: dict[str, str], key: str, default: float = 0.0) -> float:
    try:
        value = row.get(key, "")
        if value in ("", None):
            return default
        return float(value)
    except ValueError:
        return default


def safe_series(rows: Iterable[dict[str, str]], key: str) -> list[float]:
    values: list[float] = []
    for row in rows:
        value = fval(row, key, math.nan)
        if not math.isnan(value) and value != -999.999:
            values.append(value)
    return values


def minmax(rows: Iterable[dict[str, str]], key: str) -> tuple[float, float]:
    values = safe_series(rows, key)
    return (min(values), max(values)) if values else (math.nan, math.nan)


def pct(value: float) -> str:
    return f"{value:.2f}%"


def sci(value: float) -> str:
    if math.isnan(value):
        return "n/a"
    if value == 0:
        return "0"
    if abs(value) < 1e-3 or abs(value) >= 1e4:
        return f"{value:.2e}"
    return f"{value:.4g}"


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: RGBColor | None = None):
    run.font.name = "Calibri"
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.first_child_found_in("w:tblCellMar")
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            margins.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths_inches: list[float]):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    total_dxa = int(sum(widths_inches) * 1440)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths_inches:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(col)
    for row in table.rows:
        for idx, width in enumerate(widths_inches):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(table)


def style_table(table, widths_inches: list[float]):
    set_table_widths(table, widths_inches)
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    set_run_font(run, size=9.2, bold=(row_idx == 0))
            if row_idx == 0:
                shade_cell(cell, HEADER_FILL)


def draw_line_chart(path: Path, title: str, x: list[float], series: list[tuple[str, list[float], str]], y_label: str):
    width, height = 1600, 900
    margin_l, margin_r, margin_t, margin_b = 150, 270, 110, 130
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.truetype("arial.ttf", 30)
    small = ImageFont.truetype("arial.ttf", 22)
    tiny = ImageFont.truetype("arial.ttf", 18)
    draw.text((margin_l, 35), title, fill=(20, 20, 20), font=font)
    x_min, x_max = min(x), max(x)
    y_values = [v for _, vals, _ in series for v in vals]
    y_min, y_max = 0, max(100, math.ceil(max(y_values) / 10) * 10)

    def px(xv: float) -> float:
        return margin_l + (xv - x_min) / (x_max - x_min) * (width - margin_l - margin_r)

    def py(yv: float) -> float:
        return height - margin_b - (yv - y_min) / (y_max - y_min) * (height - margin_t - margin_b)

    for i in range(0, 11):
        yv = y_min + (y_max - y_min) * i / 10
        yy = py(yv)
        draw.line((margin_l, yy, width - margin_r, yy), fill=(225, 230, 236), width=1)
        draw.text((62, yy - 12), f"{yv:.0f}", fill=(80, 80, 80), font=tiny)
    draw.line((margin_l, margin_t, margin_l, height - margin_b), fill=(60, 60, 60), width=2)
    draw.line((margin_l, height - margin_b, width - margin_r, height - margin_b), fill=(60, 60, 60), width=2)
    draw.text((margin_l, height - 75), "Reaction step" if max(x) > 20 else "Stage", fill=(50, 50, 50), font=small)
    draw.text((25, margin_t + 210), y_label, fill=(50, 50, 50), font=small)

    for label, vals, color_hex in series:
        rgb = tuple(int(color_hex[i:i + 2], 16) for i in (1, 3, 5))
        points = [(px(xv), py(yv)) for xv, yv in zip(x, vals)]
        if len(points) > 1:
            draw.line(points, fill=rgb, width=5)
        for p in points[:: max(1, len(points) // 14)]:
            draw.ellipse((p[0] - 5, p[1] - 5, p[0] + 5, p[1] + 5), fill=rgb)
    legend_x = width - margin_r + 40
    legend_y = margin_t + 10
    for idx, (label, _, color_hex) in enumerate(series):
        rgb = tuple(int(color_hex[i:i + 2], 16) for i in (1, 3, 5))
        y = legend_y + idx * 38
        draw.line((legend_x, y + 12, legend_x + 45, y + 12), fill=rgb, width=5)
        draw.text((legend_x + 58, y), label, fill=(45, 45, 45), font=tiny)
    img.save(path)


def draw_transport_chart(path: Path, stage_summary: list[dict[str, float]]):
    x = [row["stage"] for row in stage_summary]
    series = [
        ("Chalcopyrite", [row["ccp"] for row in stage_summary], "#D95F02"),
        ("Bornite", [row["bn"] for row in stage_summary], "#1B9E77"),
        ("Native Au", [row["au"] for row in stage_summary], "#7570B3"),
        ("Magnetite", [row["mag"] for row in stage_summary], "#666666"),
    ]
    draw_line_chart(path, "Transport final-stage maxima by stage", x, series, "Maximum phase percentage")


def configure_doc(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("CFAS hydrothermal modeling draft")
    set_run_font(run, size=9, color=MUTED)


def add_para(doc: Document, text: str, style: str | None = None, bold_head: str | None = None):
    p = doc.add_paragraph(style=style)
    if bold_head and text.startswith(bold_head):
        r = p.add_run(bold_head)
        set_run_font(r, bold=True)
        r2 = p.add_run(text[len(bold_head):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run)


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run_font(run, size=9, color=MUTED)


def add_table(doc: Document, rows: list[list[str]], widths: list[float]):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx, text in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(text)
            set_run_font(run, size=9.2, bold=(r_idx == 0))
    style_table(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def make_doc():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    transport_cfg = json.loads((ROOT / "Project/transport/TRANSPORT.json").read_text(encoding="utf-8"))
    titration_cfg = json.loads((ROOT / "Project/titration/TITRATION.json").read_text(encoding="utf-8"))
    transport_rows = read_rows(ROOT / "Project/transport/selected_output.csv")
    titration_rows_all = read_rows(ROOT / "Project/titration/selected_output.csv")
    titration_rows = [r for r in titration_rows_all if r.get("state") == "react"]
    transport_valid = [r for r in transport_rows if fval(r, "step", -99) >= 0 and fval(r, "cell", 0) >= 1]

    titration_steps = [fval(r, "rxn_step") for r in titration_rows]
    titration_fig = ASSET_DIR / "titration_mineral_evolution.png"
    draw_line_chart(
        titration_fig,
        "Titration mineral evolution",
        titration_steps,
        [
            ("Pyrite", [fval(r, "Pyrite_pct") for r in titration_rows], "#E7298A"),
            ("Chalcopyrite", [fval(r, "Chalcopyrite(alpha)_pct") for r in titration_rows], "#D95F02"),
            ("Bornite", [fval(r, "Bornite(alpha)_pct") for r in titration_rows], "#1B9E77"),
            ("Native Au", [fval(r, "Au(element)_pct") for r in titration_rows], "#7570B3"),
        ],
        "Phase percentage",
    )

    stage_summary = []
    for stage in sorted({int(fval(r, "stage")) for r in transport_valid if fval(r, "stage", math.nan) == fval(r, "stage", math.nan)}):
        rows = [r for r in transport_valid if int(fval(r, "stage")) == stage]
        max_step = max(fval(r, "u_step") for r in rows)
        final = [r for r in rows if fval(r, "u_step") == max_step]
        stage_summary.append(
            {
                "stage": stage,
                "pH_min": min(fval(r, "pH") for r in final),
                "pH_max": max(fval(r, "pH") for r in final),
                "ccp": max(fval(r, "Chalcopyrite(alpha)_pct") for r in final),
                "bn": max(fval(r, "Bornite(alpha)_pct") for r in final),
                "au": max(fval(r, "Au(element)_pct") for r in final),
                "mag": max(fval(r, "Magnetite_pct") for r in final),
            }
        )
    transport_fig = ASSET_DIR / "transport_stage_maxima.png"
    draw_transport_chart(transport_fig, stage_summary)

    doc = Document()
    configure_doc(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(24)
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("赤铁矿参与的 Cu-Fe-S-Au 热液反应过程模拟")
    set_run_font(run, size=22, bold=True, color=RGBColor(0, 0, 0))
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("基于 PhreeqPyne 项目现有 transport 与 titration 输出的论文初稿")
    set_run_font(run, size=12, color=MUTED)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("模型版本：PhreeqPyne 0.1.0；PHREEQC/IPhreeqc 3.8.6-17100；数据库：Thermoddem V1.10 (15 Dec 2020)")
    set_run_font(run, size=9.5, color=MUTED)

    doc.add_heading("摘要", level=1)
    add_para(
        doc,
        "本文依据 PhreeqPyne 项目中已固化的 PHREEQC 输入、配置文件和 selected output，撰写一个面向地学论文的计算模拟初稿。模型用于检验在 240 °C、27 bar、弱碱性、含 Cu-Fe-S-Au-Cl 的热液体系中，赤铁矿供给、硫化物沉淀和金络合物演化是否能够在地质上可接受的条件下形成黄铁矿、黄铜矿、斑铜矿及自然金共存或分带的结果。现有结果表明，批反应滴定路径中黄铁矿在早期占优，随后被黄铜矿和斑铜矿替代，自然金比例从约 3.9% 缓慢升高到约 5.8%；一维扩散运移模型显示黄铜矿在多数阶段可成为主要沉淀相，后期磁铁矿比例增强，而斑铜矿和自然金峰值随阶段推进降低。由于项目中尚缺少样品实测约束、系统参数矩阵和失败情景集合，本文将这些结果解释为可行过程窗口的初步限定，而非对天然体系的唯一复原。"
    )
    add_para(doc, "关键词：PHREEQC；PhreeqPy；热液；黄铜矿；斑铜矿；自然金；反应路径模拟")

    doc.add_heading("1 引言", level=1)
    add_para(
        doc,
        "Cu-Fe-S-Au 热液体系中的矿物组合通常受流体组成、硫逸度、氧化还原状态、温压条件和水岩反应进程共同控制。对于含铁氧化物参与的体系，赤铁矿的溶解或被还原消耗可能改变 Fe 供给、pH 缓冲和硫化物饱和状态，从而影响黄铁矿、黄铜矿、斑铜矿及自然金的沉淀顺序。单纯列出热力学饱和指数或软件输出难以回答关键地质问题：在给定温压和流体组成下，这些矿物组合是否可以由连续水岩反应或边界流体扩散所产生，哪些条件会导致模型结果偏离目标矿物组合。"
    )
    add_para(
        doc,
        "本稿以项目当前的 transport 与 titration 两类工作流为基础，围绕“观察约束 - 模型目的 - 输入参数 - 控制方程 - 参数测试 - 允许条件窗口 - 地质对应 - 局限性”的逻辑组织内容。需要强调的是，现有仓库尚未提供样品产状、显微结构、矿物化学或流体包裹体数据，因此本文中的“观察约束”主要来自模型设计中被追踪的目标矿物相和元素种类，属于论文写作框架中的待替换约束。"
    )

    doc.add_heading("2 项目资料与模型输出基础", level=1)
    add_para(
        doc,
        "项目包含可序列化 JSON 配置、PHREEQC 输入文件、selected_output.csv 输出、Python 工作流构建器和图形界面雏形。transport 工作流模拟 10 个一维单元、8 个阶段的扩散替换过程；titration 工作流模拟单一反应溶液中相和溶液组分的逐步加入。两类模型共同追踪水相 pH、pe、Cu-Fe-S-Au 总量、含硫和含氯络合物活度、矿物相量、饱和指数及赤铁矿动力学反应物余量。"
    )
    add_table(
        doc,
        [
            ["资料来源", "项目路径", "本文使用方式"],
            ["模型配置", "Project/transport/TRANSPORT.json；Project/titration/TITRATION.json", "提取温压、初始流体、边界流体、矿物相、反应步和输出指标"],
            ["PHREEQC 脚本", "Project/*/phreeqpyne_input.phr", "核对 SOLUTION、EQUILIBRIUM_PHASES、TRANSPORT、RATES 和 SELECTED_OUTPUT 设置"],
            ["数值结果", "Project/*/selected_output.csv", "统计矿物相百分比、pH、pe、金络合物活度和阶段/反应步演化"],
            ["程序实现", "src/phreeqpyne", "确认模型生成逻辑、动力学表达式、插值方式和工作流差异"],
        ],
        [1.35, 2.65, 2.5],
    )

    doc.add_heading("3 方法", level=1)
    doc.add_heading("3.1 Numerical and geochemical modeling", level=2)
    add_para(
        doc,
        "数值模拟采用 PhreeqPyne 生成 PHREEQC/IPhreeqc 输入文件，并以 Thermoddem V1.10 热力学数据库计算高温热液条件下水相物种分配、矿物饱和和相平衡。模型服务的地质问题是：在 240 °C、27 bar、弱碱性含盐流体中，赤铁矿参与的反应路径是否能够产生 Cu-Fe 硫化物与自然金的共存组合，并限定这种组合出现的水化学和反应进程条件。当前模型主要测试边界流体输入、赤铁矿动力学消耗、逐步加入的 H2S、Cu、Au 和黄铜矿/赤铁矿组分，以及扩散替换阶段对矿物组合的影响。输出结果通过矿物相百分比、饱和指数、金硫络合物活度和 pH-pe 演化与目标矿物组合进行对比。"
    )

    doc.add_heading("3.1.1 Modeling rationale and conceptual framework", level=3)
    add_para(
        doc,
        "模型被设计为检验目标矿物组合是否能在地质上合理的高温热液条件下被重现，而不是复原唯一的天然反应路径。transport 模型代表半开放的一维扩散替换系统：外部边界流体逐阶段作用于具有壳-核浓度梯度的初始孔隙流体，矿物相允许在平衡约束下沉淀，赤铁矿以动力学反应物参与。titration 模型代表批反应/反应路径系统：单一热液溶液中逐步加入赤铁矿、黄铜矿、H2S、Cu(I) 和 Au(III)，用于测试反应进程对矿物组合转变的影响。"
    )
    add_para(
        doc,
        "模型简化了天然体系中的流体多期脉动、非均质孔隙结构、压力波动、真实矿物表面变化和复杂氧化还原缓冲。现阶段这些简化是可接受的，因为项目尚处于从 notebook 原型迁移到可复用工作流的阶段，当前目标是识别可能的反应窗口和不合理情景，而不是建立具有唯一性的矿床成因模型。"
    )

    doc.add_heading("3.1.2 Input constraints and initial conditions", level=3)
    add_para(
        doc,
        "模型输入主要来自项目配置文件，而非独立样品数据。transport 与 titration 均设置温度 240 °C、压力 27 bar、pH 8.0；transport 的边界流体水量为 0.008 kg，titration 反应溶液水量为 1.0 kg。基础流体含 Na 1.05 mol/kgw、Cl 0.5 mol/kgw、K 0.001 mol/kgw 和 P(V) 0.1 mol/kgw。transport 的边界流体中 Cu(I) 从 6.5e-4 降至 4.1e-4 mol/kgw、Fe(III) 从 3.0e-4 升至 5.0e-4 mol/kgw、S(-II) 从 2.4e-2 降至 2.0e-2 mol/kgw，Au(III) 保持 1.0e-4 mol/kgw。"
    )
    add_para(
        doc,
        "transport 初始孔隙流体包含 Na 0.1、Cl 0.1、Cu(II) 5.0e-5、Fe(III) 1.0e-5 mol/kgw，S(-II) 和 Au(III) 初值为 0。初始孔隙流体对 Na、Cl、K、P、Cu、Fe、S、Au 设置从壳部到核部降低到 0.1 倍的对数梯度。titration 模型在 100 个反应步中加入 Hematite 总量 0.001、Chalcopyrite(alpha) 总量 0.001、H2S(g) 总量 0.1、Cu(I) 总量 0.006 和 Au(III) 总量 1.0e-4，其中多数采用指数权重分配。上述数值当前应视为建模假设或工作参数，仍需由实测和文献约束替换或校准。"
    )

    doc.add_heading("3.1.3 Governing reactions/equations and thermodynamic database", level=3)
    add_para(
        doc,
        "计算使用 IPhreeqcCOM 3.8.6-17100 和 PHREEQC_ThermoddemV1.10_15Dec2020.dat。transport 过程由 PHREEQC 的 TRANSPORT 模块处理，其概念方程可写为一维扩散-反应形式：∂C_i/∂t = D∂²C_i/∂x² + R_i，其中 C_i 为水相组分浓度，D 为扩散系数，R_i 为平衡相沉淀/溶解及动力学反应项。项目配置采用 diffusion_only、constant closed 边界、单元长度 0.001 m、弥散度 1e-6 m、扩散系数 1e-9 m²/s 和 10000 s 时间步。"
    )
    add_para(
        doc,
        "赤铁矿动力学表达式在 PHREEQC RATES 块中实现。反应仅在 SI(Hematite) < 0 且反应物余量 M > 0 时进行，速率近似为 rate = A(kH + kH2O + kOH)(1 - 10^SI)，其中 A 与比表面积、粗糙度和 M/M0 有关，kH、kH2O、kOH 分别表示酸促进、中性水和 OH 相关项。该表达式在地质意义上代表赤铁矿非饱和条件下的动力学溶解或供铁过程。矿物相平衡包括 Pyrite、Chalcopyrite(alpha)、Bornite(alpha)、Magnetite、Chalcocite(alpha)、Au(element) 和 Cu(element)，titration 还使用 O2(g) 作为氧化还原/逸度缓冲项。"
    )

    doc.add_heading("3.1.4 Boundary conditions and reaction/transport setup", level=3)
    add_para(
        doc,
        "transport 模型以 SOLUTION 0 作为外部边界流体，初始 10 个单元代表壳-核梯度孔隙流体。8 个阶段的 shift 数分别为 2、3、4、10、15、16、19 和 20，输出单元为 1-10，输出频率为每个 shift。titration 模型使用 reaction_solution 作为初始溶液，逐步生成相添加和溶液组分添加，并在每一步保存 solution 和 equilibrium_phases，以便下一步继承上一反应状态。两类模型均保持温度和压力恒定。"
    )

    doc.add_heading("3.1.5 Parameter selection and sensitivity tests", level=3)
    add_para(
        doc,
        "现有仓库中真正成体系的参数敏感性测试尚未完成。已有的参数变化主要体现在 transport 的 8 阶段边界流体插值、壳-核初始梯度、titration 的 100 步反应组分权重变化，以及 Example 与 Project 配置之间的 titration 设置差异。这些设置可作为主模型和方法验证的基础，但还不足以构成完整的单因素或组合参数矩阵。后续应系统改变 pH、pe/fO2、S(-II)、Cu/Au 输入、流体/岩石比、扩散系数、赤铁矿反应量和关键平衡相容量，以限定允许条件窗口。"
    )

    doc.add_heading("3.1.6 Criteria for evaluating model consistency", level=3)
    add_para(
        doc,
        "一个情景被视为与目标约束相容，至少需要满足以下标准：能够形成黄铁矿、黄铜矿、斑铜矿或自然金中的目标组合；不会生成与目标组合明显矛盾的过量自然铜或辉铜矿；水相 pH 保持在弱碱性附近而非发生不可解释的强酸/强碱漂移；金主要以硫络合物或自然金沉淀响应硫化过程；饱和指数和相量变化具有连贯的反应顺序。无法形成黄铜矿/自然金、或仅在极端电荷不平衡和非物理输出中出现目标相的情景，应被用于排除相应参数组合。"
    )

    doc.add_heading("4 结果", level=1)
    doc.add_heading("4.1 Baseline model evolution", level=2)
    add_para(
        doc,
        "在 titration 基准模型中，系统经历三个阶段。第 1-23 步为黄铁矿-黄铜矿转换阶段：黄铁矿从 70.16% 下降到接近 0，黄铜矿从 25.92% 升高到 96.06%。第 24 步以后斑铜矿开始出现，并与黄铜矿共同控制 Cu-Fe-S 矿物组合。第 50-100 步中，黄铜矿由 74.42% 下降到 45.76%，斑铜矿由 20.88% 升高到 48.47%，自然金由约 4.70% 升高到 5.77%。整个反应路径中 pH 从 8.15 升至 8.29，pe 保持在约 -7.90 至 -7.58 的还原范围。"
    )
    add_para(
        doc,
        "transport 基准模型在 8 个阶段中显示空间扩散和阶段性边界流体共同控制的演化。阶段末端的黄铜矿最大比例由 72.83% 升至 99.45%，斑铜矿最大比例由第 2 阶段的 20.52% 逐步降低至第 8 阶段的 6.32%，自然金最大比例由第 2 阶段的 24.04% 降至第 8 阶段的 19.64%。磁铁矿在后期增强，第 8 阶段最大比例达 67.73%。"
    )
    doc.add_picture(str(titration_fig), width=Inches(6.1))
    add_caption(doc, "图 1. titration 模型中主要矿物相百分比随反应步的变化。")
    doc.add_picture(str(transport_fig), width=Inches(6.1))
    add_caption(doc, "图 2. transport 模型中各阶段末端主要矿物相最大百分比。")

    doc.add_heading("4.2 Effects of key parameters", level=2)
    add_para(
        doc,
        "反应进程对矿物组合具有一级控制作用。随着 titration 中黄铜矿、H2S、Cu(I) 和 Au(III) 的逐步加入，早期黄铁矿稳定性降低，黄铜矿成为主要沉淀相；继续反应后，体系由黄铜矿主导转向黄铜矿-斑铜矿共存。"
    )
    add_para(
        doc,
        "硫化程度和金络合物活度影响自然金响应。titration 中 la_Au(HS)2- 位于 -5.21 至 -4.89，la_AuHS 位于 -7.66 至 -7.34；自然金比例随反应进程缓慢增加，显示在还原硫化条件下金的硫络合物与自然金沉淀之间存在耦合。"
    )
    add_para(
        doc,
        "空间扩散和阶段性边界输入调节矿物空间分布。transport 模型中，pH 在 7.96-8.09 之间变化，黄铜矿在后续阶段保持高比例，而磁铁矿后期增强，说明边界流体的 Fe 输入、赤铁矿反应余量和扩散替换进程共同影响 Fe 氧化物/硫化物平衡。"
    )

    doc.add_heading("4.3 Sensitivity analysis and permissible condition window", level=2)
    add_para(
        doc,
        "基于现有单一主模型输出，可初步限定的允许窗口为：温度 240 °C、压力 27 bar、pH 约 8.0-8.3、还原 pe 条件、较高 S(-II) 供给以及 Cu-Au 同步输入。在该窗口中，模型能够形成黄铁矿-黄铜矿-斑铜矿-自然金的演化序列；黄铜矿可作为主要 Cu-Fe 硫化物相，自然金保持少量但持续出现。该窗口目前不是严格敏感性分析结果，而是由现有有效模拟限定的工作窗口。"
    )
    add_table(
        doc,
        [
            ["指标", "titration 结果范围", "transport 结果范围", "解释"],
            ["pH", "8.1507-8.2921", "7.9624-8.0888", "弱碱性条件整体保持"],
            ["黄铜矿比例", "25.92%-96.06%", "0%-99.45%", "多数有效阶段中的主导 Cu-Fe 硫化物"],
            ["斑铜矿比例", "0%-48.47%", "0%-20.52%", "在反应后期或局部条件下增强"],
            ["自然金比例", "3.85%-5.77%", "0%-24.04%", "与硫化还原过程耦合，但强度依赖模型类型"],
            ["磁铁矿比例", "0%", "0%-67.73%", "transport 后期 Fe 氧化物响应更明显"],
        ],
        [1.25, 1.55, 1.55, 2.15],
    )

    doc.add_heading("4.4 Inconsistent scenarios and excluded conditions", level=2)
    add_para(
        doc,
        "现有仓库尚未保存系统性的失败情景，因此无法严格排除完整参数空间。不过从主模型内部可以识别若干潜在不相容结果：titration 的早期步骤主要形成黄铁矿而缺少斑铜矿，不足以解释斑铜矿发育的组合；transport 的若干阶段自然金峰值下降且局部磁铁矿增强，若实际观察要求持续增强的自然金而不含明显磁铁矿，则这些阶段需要被视为不相容或仅代表局部环境。缺少低硫、低 Cu、不同 pe/fO2 和不同 pH 的对照模拟，是目前排除替代解释的主要限制。"
    )

    doc.add_heading("5 讨论", level=1)
    doc.add_heading("5.1 Model-data comparison", level=2)
    add_para(
        doc,
        "模型结果可为黄铁矿-黄铜矿-斑铜矿-自然金共生或分带提供一个过程性解释：早期 Fe-S 相优先形成，随后随着 Cu 和还原硫供给增强，黄铜矿成为主导相；进一步反应进程使斑铜矿比例增加，同时少量自然金得以沉淀。若后续样品中观察到黄铁矿被黄铜矿/斑铜矿切穿、包裹或替代，以及自然金与 Cu-Fe 硫化物共生，这一模型路径将具有较强解释力。"
    )
    add_para(
        doc,
        "目前仍不能被模型完全解释的是天然尺度上的结构约束、矿物世代关系、元素空间分布和绝对质量守恒。项目输出提供了相对百分比和水相活度，但缺少与薄片、矿物化学、LA-ICP-MS 或流体包裹体温盐度数据的直接对比。"
    )

    doc.add_heading("5.2 Main controlling factors", level=2)
    add_para(
        doc,
        "模拟表明，反应进程及 Cu-S 输入是黄铁矿向黄铜矿/斑铜矿转变的一级控制因素；Au-S 络合物活度和氧化还原状态则调节自然金是否持续出现。transport 模型中，边界流体阶段性变化和扩散距离成为二级控制因素，它们改变不同单元中的矿物比例和磁铁矿响应。控制因素可能随反应阶段转换：早期主要受 Fe-S 饱和控制，中晚期逐渐转向 Cu-S 供给和局部氧化还原缓冲控制。"
    )

    doc.add_heading("5.3 Implications for geological process", level=2)
    add_para(
        doc,
        "这些结果提示，若天然体系具有弱碱性、高温、还原且富硫的局部反应环境，赤铁矿参与的水岩反应可以为 Cu-Fe 硫化物和自然金沉淀提供可行路径。模型条件更适合作为局部反应环境的约束，而不应直接外推为区域成矿流体或矿床尺度通量。若后续观察证实硫化物具有由黄铁矿到黄铜矿/斑铜矿的替代序列，则当前模型可用于限定硫化程度、Cu/Au 输入和赤铁矿反应量的允许范围。"
    )

    doc.add_heading("5.4 Comparison with alternative models", level=2)
    add_para(
        doc,
        "替代解释可能包括简单冷却沉淀、单次流体混合、直接硫化物饱和或外来金颗粒机械加入。当前模型与赤铁矿参与的硫化还原反应模型相容，但进一步指出仅有高温含盐流体并不足够，还需要足够的还原硫和 Cu-Au 输入才能产生黄铜矿-斑铜矿-自然金组合。对于只涉及简单冷却或单次平衡的模型，若不能重现黄铁矿到黄铜矿/斑铜矿的演化顺序及自然金伴生，则其适用范围可能受限。不过，在缺少低温、混合和冷却端元对照模拟之前，本文不能排除这些替代模型在其他局部条件下成立。"
    )

    doc.add_heading("5.5 Model limitations and uncertainty", level=2)
    add_para(
        doc,
        "模型旨在识别可行过程窗口，而不是给出唯一自然路径。主要不确定性包括：Thermoddem 数据库在高温含 Cu-Au-S-Cl 络合体系中的物种和矿物热力学覆盖范围；赤铁矿动力学参数、比表面积和粗糙度取值的不确定性；初始孔隙流体和边界流体组成缺少实测约束；天然体系开放性、非均质性和多期脉动未被完整表示；模型尚未纳入完整参数矩阵和失败情景数据库。因此，当前推断应被理解为由已有配置和输出约束的允许范围。"
    )

    doc.add_heading("6 仍缺少的工作", level=1)
    missing_items = [
        "补充真实观察约束：矿物组合、显微结构、矿物先后关系、元素面扫/线扫、流体包裹体温盐度、围岩/矿石全岩或微区化学。",
        "区分参数来源：将实测值、文献值、重建值和假设值分别标注，并为每类参数建立引用或数据表。",
        "建立敏感性矩阵：系统改变 pH、pe/fO2、S(-II)、Cu/Au、Fe、流体/岩石比、扩散系数、赤铁矿量和矿物相容量。",
        "保存并解释失败情景：低硫、低 Cu、无赤铁矿、不同氧逸度、不同 Au 输入和不同边界条件下哪些模型无法重现目标组合。",
        "校验热力学数据库：确认 Thermoddem 对 Au-Cl-S、Cu-Cl-S 和相关硫化物/氧化物在 240 °C 条件下的适用性。",
        "完善图件和统计：输出可发表的阶段图、反应路径图、矿物相堆积图、络合物活度图和参数窗口图。",
        "补充质量守恒与电荷平衡筛查：剔除电荷误差过大或仅由中间 addition solution 产生的非目标输出行。",
        "形成论文引用体系：加入 PHREEQC、PhreeqPy、数据库、赤铁矿动力学、Au/Cu 热液络合和目标地质背景文献。",
        "给出地质尺度解释边界：明确模型代表局部反应环境、裂隙/孔隙尺度或样品尺度，而非直接代表区域成矿通量。",
    ]
    for item in missing_items:
        add_bullet(doc, item)

    doc.add_heading("7 结论", level=1)
    add_para(
        doc,
        "基于现有 PhreeqPyne 项目内容，240 °C、27 bar、弱碱性还原热液中赤铁矿参与的反应路径能够产生黄铁矿、黄铜矿、斑铜矿和自然金的可行演化序列。titration 模型强调反应进程和 Cu-S-Au 输入对矿物组合转变的控制，transport 模型显示阶段性边界流体和扩散替换可造成空间/阶段差异。当前结论的强度受限于缺少实测观察约束、系统敏感性测试和失败情景，因此更适合作为论文方法与结果框架的初稿，以及后续实证和参数测试工作的路线图。"
    )

    doc.add_heading("附录：当前基准模型关键参数", level=1)
    add_table(
        doc,
        [
            ["类别", "transport", "titration"],
            ["温度/压力", "240 °C / 27 bar", "240 °C / 27 bar"],
            ["pH / pe", "8.0 / 0.0 初始设置", "8.0 / 0.0 初始设置"],
            ["模型类型", "10 单元、8 阶段、一维 diffusion_only TRANSPORT", "100 步批反应/相与溶液组分逐步加入"],
            ["主要矿物相", "Pyrite, Chalcopyrite, Bornite, Magnetite, Chalcocite, Au, Cu", "O2(g), Pyrite, Chalcopyrite, Bornite, Magnetite, Chalcocite, Au, Cu, Hematite"],
            ["数据库", transport_cfg["runtime"]["database_path"].split("\\")[-1], titration_cfg["runtime"]["database_path"].split("\\")[-1]],
        ],
        [1.4, 2.55, 2.55],
    )

    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    make_doc()
