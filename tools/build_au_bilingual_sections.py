from __future__ import annotations

import csv
import math
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

from build_cfas_paper import (
    add_caption,
    add_para,
    add_table,
    configure_doc,
    fval,
    safe_series,
    set_run_font,
)


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs"
ASSET_DIR = OUT_DIR / "paper_assets"
EN_PATH = OUT_DIR / "Au_hydrothermal_modeling_sections_EN.docx"
CN_PATH = OUT_DIR / "Au_hydrothermal_modeling_sections_CN.docx"


def fonts(lang: str):
    regular = "C:/Windows/Fonts/msyh.ttc" if lang == "cn" else "arial.ttf"
    return (
        ImageFont.truetype(regular, 30),
        ImageFont.truetype(regular, 22),
        ImageFont.truetype(regular, 18),
    )


def read_rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", newline="", encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


def fmt_range(rows: list[dict[str, str]], key: str, digits: int = 3) -> str:
    vals = safe_series(rows, key)
    if not vals:
        return "n/a"
    return f"{min(vals):.{digits}g}-{max(vals):.{digits}g}"


def draw_dual_au_chart(path: Path, rows: list[dict[str, str]], lang: str):
    width, height = 1600, 900
    ml, mr, mt, mb = 150, 280, 110, 130
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font, small, tiny = fonts(lang)
    x = [fval(r, "rxn_step") for r in rows]
    au_solid = [fval(r, "Au(element)_pct") for r in rows]
    au_hs2 = [fval(r, "la_Au(HS)2-") for r in rows]
    au_hs = [fval(r, "la_AuHS") for r in rows]
    x_min, x_max = min(x), max(x)
    y1_min, y1_max = min(min(au_hs2), min(au_hs)) - 0.1, max(max(au_hs2), max(au_hs)) + 0.1
    y2_min, y2_max = 0, max(8, math.ceil(max(au_solid)))

    def px(v):
        return ml + (v - x_min) / (x_max - x_min) * (width - ml - mr)

    def py1(v):
        return height - mb - (v - y1_min) / (y1_max - y1_min) * (height - mt - mb)

    def py2(v):
        return height - mb - (v - y2_min) / (y2_max - y2_min) * (height - mt - mb)

    title = "Au aqueous species and native Au during titration" if lang == "en" else "滴定路径中 Au 水相物种与自然金演化"
    draw.text((ml, 35), title, fill=(20, 20, 20), font=font)
    for i in range(0, 6):
        val = y1_min + (y1_max - y1_min) * i / 5
        yy = py1(val)
        draw.line((ml, yy, width - mr, yy), fill=(225, 230, 236), width=1)
        draw.text((55, yy - 12), f"{val:.1f}", fill=(80, 80, 80), font=tiny)
        val2 = y2_min + (y2_max - y2_min) * i / 5
        draw.text((width - mr + 8, yy - 12), f"{val2:.1f}", fill=(80, 80, 80), font=tiny)
    draw.line((ml, mt, ml, height - mb), fill=(60, 60, 60), width=2)
    draw.line((width - mr, mt, width - mr, height - mb), fill=(60, 60, 60), width=2)
    draw.line((ml, height - mb, width - mr, height - mb), fill=(60, 60, 60), width=2)
    xlab = "Reaction step" if lang == "en" else "反应步"
    y1lab = "log activity" if lang == "en" else "活度对数"
    y2lab = "native Au (%)" if lang == "en" else "自然金 (%)"
    draw.text((ml, height - 75), xlab, fill=(50, 50, 50), font=small)
    draw.text((25, mt + 230), y1lab, fill=(50, 50, 50), font=small)
    draw.text((width - mr + 55, mt + 230), y2lab, fill=(50, 50, 50), font=small)

    def line(vals, mapper, color):
        pts = [(px(a), mapper(b)) for a, b in zip(x, vals)]
        draw.line(pts, fill=color, width=5)
        for p in pts[::10]:
            draw.ellipse((p[0] - 5, p[1] - 5, p[0] + 5, p[1] + 5), fill=color)

    line(au_hs2, py1, (31, 119, 180))
    line(au_hs, py1, (214, 39, 40))
    line(au_solid, py2, (117, 112, 179))
    legend = [("Au(HS)2-", (31, 119, 180)), ("AuHS", (214, 39, 40)), ("Au(element)", (117, 112, 179))]
    lx, ly = width - mr + 70, mt + 10
    for idx, (label, color) in enumerate(legend):
        yy = ly + idx * 40
        draw.line((lx, yy + 12, lx + 45, yy + 12), fill=color, width=5)
        draw.text((lx + 55, yy), label, fill=(45, 45, 45), font=tiny)
    img.save(path)


def draw_transport_au_chart(path: Path, rows: list[dict[str, str]], lang: str):
    width, height = 1600, 900
    ml, mr, mt, mb = 150, 260, 110, 130
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font, small, tiny = fonts(lang)
    stage_summary = []
    for stage in sorted({int(fval(r, "stage")) for r in rows if fval(r, "stage", math.nan) == fval(r, "stage", math.nan)}):
        stage_rows = [r for r in rows if int(fval(r, "stage")) == stage]
        max_step = max(fval(r, "u_step") for r in stage_rows)
        final = [r for r in stage_rows if fval(r, "u_step") == max_step]
        stage_summary.append(
            (
                stage,
                max(fval(r, "Au(element)_pct") for r in final),
                min(fval(r, "Au(element)_pct") for r in final),
            )
        )
    x = [s[0] for s in stage_summary]
    y_max = [s[1] for s in stage_summary]
    y_min = [s[2] for s in stage_summary]
    x_min, x_max = min(x), max(x)
    y0, y1 = 0, max(25, math.ceil(max(y_max)))

    def px(v):
        return ml + (v - x_min) / (x_max - x_min) * (width - ml - mr)

    def py(v):
        return height - mb - (v - y0) / (y1 - y0) * (height - mt - mb)

    title = "Native Au redistribution at transport stage ends" if lang == "en" else "Transport 阶段末端自然金再分配"
    draw.text((ml, 35), title, fill=(20, 20, 20), font=font)
    for i in range(0, 6):
        val = y0 + (y1 - y0) * i / 5
        yy = py(val)
        draw.line((ml, yy, width - mr, yy), fill=(225, 230, 236), width=1)
        draw.text((65, yy - 12), f"{val:.0f}", fill=(80, 80, 80), font=tiny)
    draw.line((ml, mt, ml, height - mb), fill=(60, 60, 60), width=2)
    draw.line((ml, height - mb, width - mr, height - mb), fill=(60, 60, 60), width=2)
    draw.text((ml, height - 75), "Stage" if lang == "en" else "阶段", fill=(50, 50, 50), font=small)
    draw.text((25, mt + 220), "native Au (%)" if lang == "en" else "自然金 (%)", fill=(50, 50, 50), font=small)
    max_pts = [(px(a), py(b)) for a, b in zip(x, y_max)]
    min_pts = [(px(a), py(b)) for a, b in zip(x, y_min)]
    draw.line(max_pts, fill=(117, 112, 179), width=5)
    draw.line(min_pts, fill=(190, 190, 190), width=4)
    for p in max_pts:
        draw.ellipse((p[0] - 6, p[1] - 6, p[0] + 6, p[1] + 6), fill=(117, 112, 179))
    for p in min_pts:
        draw.rectangle((p[0] - 5, p[1] - 5, p[0] + 5, p[1] + 5), fill=(110, 110, 110))
    lx, ly = width - mr + 50, mt + 10
    labels = [("stage maximum", (117, 112, 179)), ("stage minimum", (110, 110, 110))]
    if lang == "cn":
        labels = [("阶段最大值", (117, 112, 179)), ("阶段最小值", (110, 110, 110))]
    for idx, (label, color) in enumerate(labels):
        yy = ly + idx * 40
        draw.line((lx, yy + 12, lx + 45, yy + 12), fill=color, width=5)
        draw.text((lx + 55, yy), label, fill=(45, 45, 45), font=tiny)
    img.save(path)


def start_doc(title: str, subtitle: str) -> Document:
    doc = Document()
    configure_doc(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(title)
    set_run_font(run, size=21, bold=True, color=RGBColor(0, 0, 0))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    run = p.add_run(subtitle)
    set_run_font(run, size=11, color=RGBColor(95, 95, 95))
    return doc


def add_methods_en(doc: Document):
    doc.add_heading("3 Methods", level=1)
    doc.add_heading("3.1 Numerical and geochemical modeling", level=2)
    add_para(doc, "The primary modelling question is not only whether hematite-involved reactions in a Cu-Fe-S-Au-Cl-bearing hydrothermal fluid can generate a coherent pyrite-chalcopyrite-bornite assemblage, but also how Au-bearing aqueous species and Au-related solid phases evolve during this reaction path. Specifically, the model evaluates how Au is transported, released, precipitated, and spatially redistributed during surface-product dissolution-reprecipitation and infiltration-controlled shell-core alteration. The simulations were implemented with PhreeqPyne as a reproducible Python wrapper for PHREEQC/IPhreeqc input generation and selected-output parsing.")
    doc.add_heading("3.1.1 Modeling rationale and conceptual framework", level=3)
    add_para(doc, "The model was designed to test whether a redox- and sulfur-buffered hydrothermal reaction path can reproduce coupled Cu-Fe sulfide evolution and Au transfer under geologically plausible conditions. Two complementary conceptual systems were used. The titration model represents a batch reaction path in which phase and solution additions drive progressive dissolution-reprecipitation. The transport model represents a semi-open one-dimensional infiltration system in which an external boundary fluid modifies an initially zoned shell-core pore fluid by diffusion-only replacement.")
    add_para(doc, "The model focuses on phase transfer among hematite, pyrite, chalcopyrite, bornite, magnetite, and native Au, together with Au-bearing aqueous complexes such as Au(HS)2- and AuHS. It simplifies the natural system by excluding explicitly resolved fracture networks, variable permeability, episodic boiling, mechanical remobilization of native Au, and independently changing temperature-pressure histories. These simplifications are acceptable here because the purpose is to delimit feasible reaction windows rather than to reconstruct a unique natural pathway.")
    doc.add_heading("3.1.2 Input constraints and initial conditions", level=3)
    add_para(doc, "The calculations use the project configuration files as the current input constraints. Both model families are set at 240 degrees C and 27 bar with an initial pH of 8.0. The transport boundary fluid contains Na, Cl, K, P(V), Cu(I), Fe(III), S(-II), and Au(III); Cu(I) decreases from 6.5e-4 to 4.1e-4 mol/kgw, Fe(III) increases from 3.0e-4 to 5.0e-4 mol/kgw, S(-II) decreases from 2.4e-2 to 2.0e-2 mol/kgw, and Au(III) is held at 1.0e-4 mol/kgw across the staged boundary fluids.")
    add_para(doc, "The initial transport pore fluid is lower in salinity and metal content than the boundary fluid and is assigned a shell-core gradient for Na, Cl, K, P, Cu, Fe, S, and Au. In the titration workflow, the starting solution is reacted through 100 steps with scheduled additions of hematite, chalcopyrite, H2S(g), Cu(I), and Au(III). These settings represent the present modelling scenario and are treated as working constraints for evaluating Au transport and precipitation behavior.")
    doc.add_heading("3.1.3 Governing reactions/equations and thermodynamic database", level=3)
    add_para(doc, "The simulations use IPhreeqcCOM 3.8.6-17100 with the PHREEQC_ThermoddemV1.10_15Dec2020.dat thermodynamic database. The transport calculation follows a one-dimensional diffusion-reaction form, in which aqueous components are redistributed by diffusion while mineral equilibrium and kinetic reactions provide source and sink terms. Hematite dissolution is represented by a kinetic rate law that proceeds only when SI(Hematite) is negative and remaining kinetic reactant is available.")
    add_para(doc, "The key geochemical pathways represented by the model are hematite involvement, sulfide precipitation, Au complex transport, native-Au precipitation, and secondary redistribution during shell-core alteration. Au behavior is tracked through total Au, Au(I)/Au(III), Au(HS)2-, AuHS, Au-Cl species, Au hydroxide species, Au(element), and the saturation index of Au(element).")
    doc.add_heading("3.1.4 Boundary conditions and reaction/transport setup", level=3)
    add_para(doc, "The transport model consists of 10 cells and 8 reaction stages with stage shifts of 2, 3, 4, 10, 15, 16, 19, and 20. The boundary composition represents an infiltrating external hydrothermal fluid, whereas the initial cell solutions represent a shell-core pore-fluid gradient. The transport setup uses diffusion_only flow, constant closed boundary conditions, a 10000 s time step, 0.001 m cell length, 1e-6 m dispersivity, and a diffusion coefficient of 1e-9 m2/s.")
    add_para(doc, "The titration model is run as a stepwise phase-plus-solution reaction path. Each reaction step saves the evolving solution and equilibrium-phase assemblage for reuse in the next step, allowing progressive release and sequestration of Au to be evaluated along the same reaction path that produces the Cu-Fe sulfide assemblage.")
    doc.add_heading("3.1.5 Parameter selection and sensitivity tests", level=3)
    add_para(doc, "The current model explores parameter effects through staged boundary-fluid interpolation, shell-core initial gradients, and scheduled titration additions. The most important varied quantities are reaction progress, Cu(I) input, S(-II)/H2S input, Au(III) input, hematite addition, chalcopyrite addition, and the transport stage. These tests evaluate whether Au remains mobile as aqueous complexes, is released during surface-product turnover, or is trapped as native Au during sulfide precipitation.")
    doc.add_heading("3.1.6 Criteria for evaluating model consistency", level=3)
    add_para(doc, "A scenario is considered consistent with the modelling target only if it reproduces a coherent pyrite-chalcopyrite-bornite evolution while also producing interpretable Au behavior. Specifically, the output must show (1) Au-bearing aqueous complexes during transport or reaction, (2) precipitation or persistence of Au(element) without excessive unrelated native Cu or chalcocite dominance, (3) a stable weakly alkaline hydrothermal pH range, and (4) spatial or temporal Au redistribution that can be related to infiltration-controlled shell-core alteration.")


def add_results_en(doc: Document, titration_rows, transport_rows, fig1, fig2):
    doc.add_heading("4 Results", level=1)
    doc.add_heading("4.1 Baseline model evolution", level=2)
    add_para(doc, f"In the titration model, the pyrite-chalcopyrite-bornite sequence develops together with systematic Au redistribution. Native Au remains present throughout the reaction path and increases from 3.93% at step 1 to 5.77% at step 100. The dominant Au-bearing aqueous species remain Au-sulfur complexes: log aAu(HS)2- varies from {fmt_range(titration_rows, 'la_Au(HS)2-', 4)}, whereas log aAuHS varies from {fmt_range(titration_rows, 'la_AuHS', 4)}. This indicates that Au is not simply removed at the first sulfide saturation step; it remains coupled to the evolving sulfur-bearing fluid before being progressively partitioned into Au(element).")
    add_para(doc, "The Cu-Fe sulfide assemblage evolves from early pyrite dominance to chalcopyrite dominance and then toward chalcopyrite-bornite coexistence. Pyrite decreases from 70.16% at step 1 to nearly zero by step 24, chalcopyrite reaches 96.06% near step 23, and bornite rises to 48.47% by step 100. The Au response is therefore synchronous with the surface-product dissolution-reprecipitation sequence rather than independent of the sulfide reaction path.")
    doc.add_picture(str(fig1), width=Inches(6.1))
    add_caption(doc, "Figure 1. Evolution of Au-bearing aqueous species and native Au during the titration reaction path.")
    add_para(doc, "In the transport model, Au(element) is spatially redistributed during staged infiltration. The maximum native-Au percentage at stage ends is highest during stages 2-3, reaching 24.04% and 23.82%, then gradually decreases to 19.64% by stage 8. This trend indicates that Au precipitation is favored during early to intermediate infiltration-controlled alteration, whereas later stages increasingly stabilize chalcopyrite and magnetite at the expense of the highest native-Au accumulation.")
    doc.add_picture(str(fig2), width=Inches(6.1))
    add_caption(doc, "Figure 2. Stage-end native-Au redistribution in the one-dimensional transport model.")
    doc.add_heading("4.2 Effects of key parameters", level=2)
    add_para(doc, "Increasing reaction progress promotes transfer from pyrite to chalcopyrite and then to bornite, while native Au increases more gradually. Higher sulfur availability supports Au-sulfur complexes and allows Au to remain mobile along the reaction path before precipitation. In contrast, the transport stages show that Au accumulation is not monotonic; it depends on where infiltrating fluid, local sulfide precipitation, and evolving Fe-bearing phases intersect within the shell-core gradient.")
    add_para(doc, "The Au-bearing aqueous species respond more sensitively to sulfur-complexing conditions than to total Au alone. Au(HS)2- remains the higher-activity tracked Au complex relative to AuHS during the titration path, consistent with Au transport in reduced sulfur-bearing hydrothermal fluid. Native Au increases only after the sulfide assemblage has begun reorganizing, suggesting that Au precipitation is tied to reaction-front evolution and local changes in ligand availability.")
    doc.add_heading("4.3 Sensitivity analysis and permissible condition window", level=2)
    add_para(doc, "The current model constrains a working permissible window at 240 degrees C, 27 bar, weakly alkaline pH, reduced sulfur-bearing conditions, and coupled Cu-Au input. Within this window, Au can be transported as sulfur complexes, released during reaction-front advancement, and precipitated as Au(element) together with a coherent Cu-Fe sulfide assemblage. Outside this window, the expected failure modes would include lack of Au precipitation, sulfide assemblages dominated by only pyrite or only chalcopyrite, or Au redistribution that is decoupled from shell-core alteration.")
    add_table(doc, [
        ["Output constraint", "Titration result", "Transport result", "Modelling interpretation"],
        ["Au aqueous species", f"log aAu(HS)2- = {fmt_range(titration_rows, 'la_Au(HS)2-', 4)}; log aAuHS = {fmt_range(titration_rows, 'la_AuHS', 4)}", "Au complexes tracked with staged infiltration", "Au remains mobile as reduced sulfur complexes"],
        ["Native Au", "3.85-5.77%", "0-24.04%", "Au precipitation is progressive and spatially variable"],
        ["Sulfide assemblage", "pyrite -> chalcopyrite -> bornite", "chalcopyrite-rich with local bornite", "Au evolution is coupled to Cu-Fe sulfide turnover"],
        ["pH range", f"{fmt_range(titration_rows, 'pH', 4)}", f"{fmt_range(transport_rows, 'pH', 4)}", "weakly alkaline conditions remain stable"],
    ], [1.35, 1.65, 1.45, 2.05])
    doc.add_heading("4.4 Inconsistent scenarios and excluded conditions", level=2)
    add_para(doc, "The current tested path excludes interpretations in which Au behavior is independent of sulfide reaction progress. Early titration steps with high pyrite but low bornite do not represent the full Au-bearing assemblage, whereas later transport stages with decreasing native-Au maxima indicate that infiltration alone does not monotonically concentrate Au. These internally inconsistent intervals help delimit the conditions under which Au precipitation is linked to reaction-front chemistry rather than arbitrary Au input.")


def add_discussion_en(doc: Document):
    doc.add_heading("5 Discussion", level=1)
    doc.add_heading("5.1 Model-data comparison", level=2)
    add_para(doc, "The model provides a process-based explanation for Au redistribution during Cu-Fe sulfide replacement. A natural assemblage showing pyrite replaced by chalcopyrite-bornite, with native Au associated with sulfide reaction fronts, would be consistent with the modeled sequence. The coexistence of Au-sulfur complexes and Au(element) in the output supports a pathway in which Au is transported in reduced sulfur-bearing fluid and precipitated during local ligand, redox, or mineral-buffer changes.")
    doc.add_heading("5.2 Main controlling factors", level=2)
    add_para(doc, "The simulations indicate that sulfur-complexing capacity and reaction progress are primary controls on Au mobility, whereas Cu-Fe sulfide saturation controls the timing and location of Au precipitation. In the transport model, infiltration stage and shell-core gradients modulate where Au accumulates. Thus the dominant control shifts from Au transport as aqueous sulfur complexes to Au trapping during dissolution-reprecipitation at the reaction front.")
    doc.add_heading("5.3 Implications for geological process", level=2)
    add_para(doc, "The modeled conditions suggest that Au enrichment can occur locally during surface-product dissolution-reprecipitation rather than requiring a separate Au-only event. Au transport, release, and precipitation are best interpreted as part of the same hydrothermal alteration system that reorganizes pyrite, chalcopyrite, and bornite. The results apply most directly to local reaction fronts and shell-core alteration domains rather than to regional metal budgets.")
    doc.add_heading("5.4 Comparison with alternative models", level=2)
    add_para(doc, "Models invoking simple Au saturation from a homogeneous fluid are partly consistent with the presence of native Au, but they do not by themselves explain the coupled evolution of Au complexes, native Au, and Cu-Fe sulfide replacement. The present model further requires reduced sulfur-complexing conditions and an infiltration or reaction-front geometry. A purely mechanical redistribution model for native Au may still operate locally, but it is not required to explain the simulated Au-solid response.")
    doc.add_heading("5.5 Model limitations and uncertainty", level=2)
    add_para(doc, "The model identifies plausible process windows rather than a unique natural pathway. Key uncertainties include thermodynamic data for high-temperature Au-Cl-S complexes, kinetic parameters for hematite involvement, non-unique initial fluids, and the simplified representation of natural permeability and reaction surfaces. The inferred conditions should therefore be treated as permissible ranges for Au transport and precipitation during Cu-Fe sulfide alteration.")


def add_methods_cn(doc: Document):
    doc.add_heading("3 方法", level=1)
    doc.add_heading("3.1 数值与地球化学模拟", level=2)
    add_para(doc, "本研究的核心建模问题不仅是检验含赤铁矿反应的 Cu-Fe-S-Au-Cl 热液能否形成连贯的黄铁矿-黄铜矿-斑铜矿组合，还包括 Au-bearing 水相物种和 Au 相关固相在该反应路径中的演化。具体而言，模型用于评估在表面产物溶解-再沉淀和入渗控制的壳-核蚀变过程中，Au 如何被搬运、释放、沉淀和空间再分配。模拟由 PhreeqPyne 生成 PHREEQC/IPhreeqc 输入并解析 selected output。")
    doc.add_heading("3.1.1 模型依据与概念框架", level=3)
    add_para(doc, "模型用于测试在地质上合理的还原、富硫热液条件下，Cu-Fe 硫化物演化与 Au 转移是否能够同时被重现。本文使用两个互补概念系统：titration 模型代表批反应路径，通过相和溶液组分逐步加入推动溶解-再沉淀；transport 模型代表半开放一维入渗体系，外部边界流体以 diffusion_only 方式改造具有初始壳-核梯度的孔隙流体。")
    add_para(doc, "模型重点关注赤铁矿、黄铁矿、黄铜矿、斑铜矿、磁铁矿和自然金之间的相转移，以及 Au(HS)2-、AuHS 等 Au-bearing 水相络合物。模型未显式处理裂隙网络、渗透率变化、沸腾、自然金机械再搬运和独立变化的温压路径；这些简化是可接受的，因为本文目标是限定可行反应窗口，而不是复原唯一自然路径。")
    doc.add_heading("3.1.2 输入约束与初始条件", level=3)
    add_para(doc, "当前输入约束来自项目配置文件。两类模型均设定为 240 °C、27 bar、初始 pH 8.0。transport 边界流体含 Na、Cl、K、P(V)、Cu(I)、Fe(III)、S(-II) 和 Au(III)；其中 Cu(I) 由 6.5e-4 降至 4.1e-4 mol/kgw，Fe(III) 由 3.0e-4 升至 5.0e-4 mol/kgw，S(-II) 由 2.4e-2 降至 2.0e-2 mol/kgw，Au(III) 在各阶段保持 1.0e-4 mol/kgw。")
    add_para(doc, "transport 初始孔隙流体较边界流体盐度和金属含量低，并对 Na、Cl、K、P、Cu、Fe、S 和 Au 设置壳-核梯度。titration 模型在 100 个反应步中逐步加入 Hematite、Chalcopyrite(alpha)、H2S(g)、Cu(I) 和 Au(III)。这些设定代表当前建模场景，用于评价 Au 的搬运和沉淀行为。")
    doc.add_heading("3.1.3 控制反应、方程与热力学数据库", level=3)
    add_para(doc, "模拟使用 IPhreeqcCOM 3.8.6-17100 和 PHREEQC_ThermoddemV1.10_15Dec2020.dat。transport 计算可概化为一维扩散-反应形式：水相组分由扩散再分配，矿物平衡和动力学反应提供源汇项。赤铁矿动力学仅在 SI(Hematite) 为负且仍有动力学反应物余量时进行。")
    add_para(doc, "模型表示的关键地球化学过程包括赤铁矿参与、硫化物沉淀、Au 络合物搬运、自然金沉淀以及壳-核蚀变中的次生再分配。Au 行为通过总 Au、Au(I)/Au(III)、Au(HS)2-、AuHS、Au-Cl 物种、Au 羟基物种、Au(element) 及 Au(element) 饱和指数共同追踪。")
    doc.add_heading("3.1.4 边界条件与反应/运移设置", level=3)
    add_para(doc, "transport 模型包含 10 个单元和 8 个反应阶段，阶段 shift 分别为 2、3、4、10、15、16、19 和 20。边界流体代表入渗热液，初始单元溶液代表壳-核孔隙流体梯度。运移设置为 diffusion_only、constant closed 边界、10000 s 时间步、0.001 m 单元长度、1e-6 m 弥散度和 1e-9 m2/s 扩散系数。")
    add_para(doc, "titration 模型作为逐步相-溶液反应路径运行。每个反应步保存演化后的 solution 和 equilibrium_phases，并在下一步继续使用，从而在形成 Cu-Fe 硫化物组合的同一反应路径中评价 Au 的释放和固定。")
    doc.add_heading("3.1.5 参数选择与敏感性测试", level=3)
    add_para(doc, "现有模型通过阶段性边界流体插值、壳-核初始梯度和滴定反应组分排程考察参数影响。最重要的变化量包括反应进程、Cu(I) 输入、S(-II)/H2S 输入、Au(III) 输入、赤铁矿加入、黄铜矿加入和 transport 阶段。这些测试用于判断 Au 是保持为水相络合物、在表面产物转换中释放，还是在硫化物沉淀过程中固定为自然金。")
    doc.add_heading("3.1.6 模型一致性判断标准", level=3)
    add_para(doc, "合理情景必须同时重现黄铁矿-黄铜矿-斑铜矿演化和可解释的 Au 行为。具体标准包括：(1) 反应或运移过程中存在 Au-bearing 水相络合物；(2) Au(element) 能够沉淀或持续存在，且不伴随过量无关自然铜或辉铜矿主导；(3) pH 保持在弱碱性热液范围；(4) Au 的时间或空间再分配能够与入渗控制的壳-核蚀变相联系。")


def add_results_cn(doc: Document, titration_rows, transport_rows, fig1, fig2):
    doc.add_heading("4 结果", level=1)
    doc.add_heading("4.1 基准模型演化", level=2)
    add_para(doc, f"在 titration 模型中，黄铁矿-黄铜矿-斑铜矿序列与系统性的 Au 再分配同步发生。自然金在整个反应路径中持续存在，并由第 1 步的 3.93% 增加到第 100 步的 5.77%。主要 Au-bearing 水相物种为 Au-硫络合物：log aAu(HS)2- 的范围为 {fmt_range(titration_rows, 'la_Au(HS)2-', 4)}，log aAuHS 的范围为 {fmt_range(titration_rows, 'la_AuHS', 4)}。这说明 Au 并非在最初硫化物饱和时即被完全移除，而是在富硫流体中持续参与络合、迁移和逐步固相分配。")
    add_para(doc, "Cu-Fe 硫化物组合从早期黄铁矿占优，转向黄铜矿占优，并最终向黄铜矿-斑铜矿共存演化。黄铁矿从第 1 步的 70.16% 降至第 24 步附近接近 0，黄铜矿在第 23 步附近达到 96.06%，斑铜矿到第 100 步升至 48.47%。Au 的响应因此与表面产物溶解-再沉淀序列同步，而不是独立于硫化物反应路径。")
    doc.add_picture(str(fig1), width=Inches(6.1))
    add_caption(doc, "图 1. 滴定反应路径中 Au-bearing 水相物种与自然金的演化。")
    add_para(doc, "在 transport 模型中，Au(element) 在阶段性入渗过程中发生空间再分配。阶段末端自然金最大比例在第 2-3 阶段最高，分别达到 24.04% 和 23.82%，随后到第 8 阶段逐渐降低至 19.64%。该趋势表明，自然金沉淀更有利于早期至中期入渗控制的蚀变阶段，而后期阶段更倾向于稳定黄铜矿和磁铁矿，并降低最高自然金富集程度。")
    doc.add_picture(str(fig2), width=Inches(6.1))
    add_caption(doc, "图 2. 一维 transport 模型中阶段末端自然金再分配。")
    doc.add_heading("4.2 关键参数效应", level=2)
    add_para(doc, "反应进程增强会推动黄铁矿向黄铜矿再到斑铜矿转化，而自然金比例更缓慢上升。较高硫供给支持 Au-硫络合物，使 Au 在反应路径中保持迁移能力，然后逐步沉淀。transport 阶段显示 Au 富集并非单调增强，而取决于入渗流体、局部硫化物沉淀和壳-核梯度中含 Fe 相演化的交汇位置。")
    add_para(doc, "Au-bearing 水相物种对硫络合条件的响应比对总 Au 输入本身更敏感。在 titration 路径中，Au(HS)2- 活度高于 AuHS，符合还原富硫热液中 Au 搬运的特征。自然金在硫化物组合开始重组后逐渐增加，表明 Au 沉淀受反应前缘演化和局部配体可用性变化控制。")
    doc.add_heading("4.3 敏感性分析与允许条件窗口", level=2)
    add_para(doc, "现有模型限定的工作窗口为 240 °C、27 bar、弱碱性 pH、还原富硫条件和 Cu-Au 耦合输入。在这一窗口内，Au 可作为硫络合物被搬运，在反应前缘推进过程中释放，并与连贯的 Cu-Fe 硫化物组合一起以 Au(element) 形式沉淀。超出该窗口的预期失败模式包括缺少自然金沉淀、硫化物组合仅由黄铁矿或黄铜矿单一主导，或 Au 再分配与壳-核蚀变过程脱耦。")
    add_table(doc, [
        ["输出约束", "titration 结果", "transport 结果", "建模解释"],
        ["Au 水相物种", f"log aAu(HS)2- = {fmt_range(titration_rows, 'la_Au(HS)2-', 4)}；log aAuHS = {fmt_range(titration_rows, 'la_AuHS', 4)}", "随阶段性入渗被追踪", "Au 以还原硫络合物保持迁移能力"],
        ["自然金", "3.85-5.77%", "0-24.04%", "Au 沉淀具有渐进性和空间差异"],
        ["硫化物组合", "黄铁矿 -> 黄铜矿 -> 斑铜矿", "黄铜矿占优并局部含斑铜矿", "Au 演化与 Cu-Fe 硫化物转换耦合"],
        ["pH 范围", f"{fmt_range(titration_rows, 'pH', 4)}", f"{fmt_range(transport_rows, 'pH', 4)}", "弱碱性条件整体稳定"],
    ], [1.35, 1.65, 1.45, 2.05])
    doc.add_heading("4.4 不一致情景与被排除条件", level=2)
    add_para(doc, "当前测试路径排除了 Au 行为完全独立于硫化物反应进程的解释。titration 早期高黄铁矿、低斑铜矿的步骤不能代表完整 Au-bearing 组合；transport 后期自然金最大值降低则说明入渗本身不会单调富集 Au。这些模型内部的不一致区间有助于限定 Au 沉淀必须与反应前缘化学条件相联系，而不能仅归因于任意 Au 输入。")


def add_discussion_cn(doc: Document):
    doc.add_heading("5 讨论", level=1)
    doc.add_heading("5.1 模型-资料对比", level=2)
    add_para(doc, "模型为 Cu-Fe 硫化物替代过程中的 Au 再分配提供了过程性解释。如果天然组合表现为黄铁矿被黄铜矿-斑铜矿替代，且自然金与硫化物反应前缘伴生，则与本模型序列相容。输出中 Au-硫络合物与 Au(element) 共存，支持 Au 先在还原富硫流体中搬运，再在局部配体、氧化还原或矿物缓冲变化中沉淀的路径。")
    doc.add_heading("5.2 主要控制因素", level=2)
    add_para(doc, "模拟表明，硫络合能力和反应进程是 Au 迁移性的一级控制因素，而 Cu-Fe 硫化物饱和控制 Au 沉淀的时间和位置。在 transport 模型中，入渗阶段和壳-核梯度调节 Au 的富集位置。因此，主控因素从 Au 作为水相硫络合物搬运，逐步转向反应前缘溶解-再沉淀过程中的 Au 固定。")
    doc.add_heading("5.3 对地质过程的意义", level=2)
    add_para(doc, "模型条件提示，Au 富集可在表面产物溶解-再沉淀过程中局部发生，而不一定需要独立的纯 Au 事件。Au 的搬运、释放和沉淀最好被理解为重组黄铁矿、黄铜矿和斑铜矿的同一热液蚀变系统的一部分。结果最适用于局部反应前缘和壳-核蚀变域，而不是区域尺度金属通量。")
    doc.add_heading("5.4 与替代模型比较", level=2)
    add_para(doc, "均一流体简单达到 Au 饱和的模型可以解释自然金存在，但难以单独说明 Au 络合物、自然金和 Cu-Fe 硫化物替代的耦合演化。本文模型进一步要求还原富硫络合条件和入渗或反应前缘几何。自然金机械再分配仍可能局部发生，但并不是解释模拟中 Au 固相响应的必要条件。")
    doc.add_heading("5.5 模型局限性与不确定性", level=2)
    add_para(doc, "模型用于识别可行过程窗口，而非唯一自然路径。主要不确定性包括高温 Au-Cl-S 络合物热力学数据、赤铁矿参与的动力学参数、非唯一初始流体，以及天然渗透率和反应表面的简化表达。因此，推断条件应视为 Cu-Fe 硫化物蚀变过程中 Au 搬运与沉淀的允许范围。")


def build_docs():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    titration_all = read_rows(ROOT / "Project/titration/selected_output.csv")
    titration_rows = [r for r in titration_all if r.get("state") == "react"]
    transport_all = read_rows(ROOT / "Project/transport/selected_output.csv")
    transport_rows = [r for r in transport_all if fval(r, "step", -99) >= 0 and fval(r, "cell", 0) >= 1]

    fig_en_1 = ASSET_DIR / "au_titration_species_en.png"
    fig_en_2 = ASSET_DIR / "au_transport_redistribution_en.png"
    fig_cn_1 = ASSET_DIR / "au_titration_species_cn.png"
    fig_cn_2 = ASSET_DIR / "au_transport_redistribution_cn.png"
    draw_dual_au_chart(fig_en_1, titration_rows, "en")
    draw_transport_au_chart(fig_en_2, transport_rows, "en")
    draw_dual_au_chart(fig_cn_1, titration_rows, "cn")
    draw_transport_au_chart(fig_cn_2, transport_rows, "cn")

    en = start_doc(
        "Au Transport and Precipitation During Cu-Fe-S Hydrothermal Reaction Paths",
        "Methods, results, and modelling-focused discussion",
    )
    add_methods_en(en)
    add_results_en(en, titration_rows, transport_rows, fig_en_1, fig_en_2)
    add_discussion_en(en)
    en.save(EN_PATH)

    cn = start_doc(
        "Cu-Fe-S 热液反应路径中的 Au 搬运与沉淀",
        "方法、结果和与模拟相关的讨论",
    )
    add_methods_cn(cn)
    add_results_cn(cn, titration_rows, transport_rows, fig_cn_1, fig_cn_2)
    add_discussion_cn(cn)
    cn.save(CN_PATH)
    print(EN_PATH)
    print(CN_PATH)


if __name__ == "__main__":
    build_docs()
